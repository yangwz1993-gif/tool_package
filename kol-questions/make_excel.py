#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""final.json → Excel"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

src = sys.argv[1] if len(sys.argv) > 1 else "results/final.json"
out = sys.argv[2] if len(sys.argv) > 2 else "agent测试题.xlsx"
data = json.load(open(src))

wb = Workbook(); ws = wb.active
ws.title = "题目"
headers = ["题目id", "测试题目(task)", "难度", "验收标准", "来源平台", "博主", "原文链接", "时间"]
rows = [[f"Q{i:02d}", d.get("task",""), d.get("difficulty",""), d.get("acceptance",""),
         d.get("source",""), d.get("blogger",""), d.get("link",""), d.get("date","")] for i, d in enumerate(data, 1)]

hfont = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="4472C4")
thin = Side(style="thin", color="D9D9D9"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
for c, h in enumerate(headers, 1):
    cell = ws.cell(1, c, h); cell.font = hfont; cell.fill = hfill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = border
for r, row in enumerate(rows, 2):
    for c, v in enumerate(row, 1):
        cell = ws.cell(r, c, v); cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if c == 7 and v: cell.hyperlink = v; cell.font = Font(color="0563C1", underline="single")
for i, w in enumerate([10, 60, 10, 50, 12, 18, 50, 12], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:H{len(rows)+1}"
for r in range(2, len(rows) + 2): ws.row_dimensions[r].height = 90
wb.save(out)
print(f"已保存: {out} ({len(rows)} 条)")
